# -*- coding: utf-8 -*-
"""Tanger, ville-vitrine : planches d'analyse urbaine (cartes issues de données réelles)."""
import math
import pandas as pd
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

W, H = 1920, 1080
EMU = 6350
TF = "Barlow Condensed"; BF = "Barlow"
BG = "E9E7E2"; INK = "2B2B2B"; MUTED = "77736C"; RULE = "C9C4BA"; WHITE = "FFFFFF"
RED = "C8462B"; BLUE = "1F3A5F"; GREEN = "5E7F4F"; GOLD = "D9A441"; GREY = "B9B1A5"; SALMON = "DFA283"
J = "jpg/"
TOTAL = 18
T = pd.read_csv("a_table.csv").set_index("arr")
ORDER = ["Tanger-Médina", "Charf-Souani", "Charf-Mghogha", "Beni Makada"]

prs = Presentation(); prs.slide_width = Emu(W * EMU); prs.slide_height = Emu(H * EMU)
BL = prs.slide_layouts[6]


def px(v): return Emu(int(round(v * EMU)))
def rgb(h): return RGBColor.from_string(h)


def _ln(shape, color, width, dash=None):
    if color is None:
        shape.line.fill.background(); return
    shape.line.color.rgb = rgb(color); shape.line.width = px(width)
    if dash:
        ln = shape.line._get_or_add_ln(); d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", dash)


def rect(s, x, y, w, h, fill=None, line=None, lw=1.5, dash=None, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if fill: r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
    else: r.fill.background()
    _ln(r, line, lw, dash); return r


def oval(s, cx, cy, r, fill=None, line=None, lw=1.5):
    return rect(s, cx - r, cy - r, 2 * r, 2 * r, fill, line, lw, shape=MSO_SHAPE.OVAL)


def seg(s, x1, y1, x2, y2, color=INK, lw=1.5, dash=None, head=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    c.line.color.rgb = rgb(color); c.line.width = px(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", dash)
    if head:
        e = etree.SubElement(ln, qn("a:tailEnd")); e.set("type", "triangle")
    return c


def text(s, x, y, w, h, content, size=11, color=INK, bold=False, italic=False, font=BF, align="l", anchor="t", spacing=1.2, spc=0, caps=False):
    tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = content.split("\n") if isinstance(content, str) else content
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        para.line_spacing = spacing
        runs = [(p, {})] if isinstance(p, str) else p
        for t, o in runs:
            r = para.add_run(); r.text = t.upper() if o.get("caps", caps) else t
            f = r.font; f.name = o.get("font", font); f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold); f.italic = o.get("italic", italic); f.color.rgb = rgb(o.get("color", color))
            sp = o.get("spc", spc)
            if sp:
                r._r.get_or_add_rPr().set("spc", str(int(sp)))
    return tb


def pic(s, f, x, y, w=None, h=None):
    im = Image.open(J + f); iw, ih = im.size
    if w and not h: h = w * ih / iw
    if h and not w: w = h * iw / ih
    s.shapes.add_picture(J + f, px(x), px(y), px(w), px(h)); return w, h


def slide(num, title, sub=None, kpis=(), dark=False):
    s = prs.slides.add_slide(BL)
    bg = "232526" if dark else BG; fg = "EDEDED" if dark else INK; mu = "A6A29B" if dark else MUTED
    rect(s, 0, 0, W, H, fill=bg)
    text(s, 40, 36, 1300, 50, title, size=24, font=TF, color=fg, spc=300, caps=True)
    if sub:
        text(s, 40, 90, 1250, 50, sub, size=10.5, color=mu, spacing=1.25)
    kx = W - 40
    for v, l in reversed(kpis):
        w_ = max(120, 15 * len(v) + 20, 6.2 * len(l))
        kx -= w_
        text(s, kx, 34, w_, 44, v, size=24, font=TF, bold=True, color=fg, align="r")
        text(s, kx, 86, w_, 40, l, size=8.5, color=mu, align="r", spacing=1.1)
        kx -= 34
    seg(s, 40, H - 44, W - 40, H - 44, "4A4C4D" if dark else RULE, 1)
    text(s, 40, H - 34, 800, 20, "TANGER, VILLE-VITRINE  ·  ÉTUDE DE CAS", size=8, color=mu, font=TF, spc=300)
    text(s, W - 840, H - 34, 800, 20, f"[NOM]  ·  [COURS]  ·  [DATE]        {num:02d} / {TOTAL}", size=8, color=mu, font=TF, spc=300, align="r")
    return s


def ptitle(s, x, y, t, color=INK, w=900):
    text(s, x, y, w, 20, t, size=9, bold=True, color=color, font=TF, spc=250, caps=True)


def lecture(s, x, y, w, body, title="Lecture", color=INK, size=10, h=300):
    ptitle(s, x, y, title)
    text(s, x, y + 26, w, h, body, size=size, color=color, spacing=1.35)


def legend(s, x, y, items, w=520, dark=False):
    fg = "E0E0E0" if dark else INK
    for i, (kind, c, lab) in enumerate(items):
        yy = y + i * 28
        if kind == "box": rect(s, x, yy + 3, 26, 16, fill=c)
        elif kind == "hatch": rect(s, x, yy + 3, 26, 16, fill=None, line=c, lw=2)
        elif kind == "line": seg(s, x, yy + 11, x + 26, yy + 11, c, 3.5)
        elif kind == "dash": seg(s, x, yy + 11, x + 26, yy + 11, c, 2.2, dash="dash")
        elif kind == "dot": oval(s, x + 13, yy + 11, 7, fill=c, line=WHITE, lw=1)
        elif kind == "diamond": rect(s, x + 6, yy + 4, 14, 14, fill=c, shape=MSO_SHAPE.DIAMOND)
        elif kind == "sq": rect(s, x + 7, yy + 5, 12, 12, fill=c)
        text(s, x + 40, yy + 1, w - 40, 24, lab, size=9.5, color=fg)
    return y + len(items) * 28


def badge(s, cx, cy, n, c=INK, r=13, size=9):
    oval(s, cx, cy, r, fill=c, line=WHITE, lw=1.5)
    text(s, cx - r, cy - r, 2 * r, 2 * r, str(n), size=size, bold=True, color=WHITE, align="c", anchor="m", font=TF)


def hbars(s, x, y, w, title, key, fmt, unit="", hl="Beni Makada", note=None, maxv=None):
    ptitle(s, x, y, title)
    vals = [T.loc[a, key] for a in ORDER]; mv = maxv or max(vals)
    for i, a in enumerate(ORDER):
        yy = y + 30 + i * 30; v = T.loc[a, key]
        c = RED if a == hl else GREY
        text(s, x, yy, 120, 22, a, size=9, color=INK, bold=(a == hl))
        bw = max(2, (w - 190) * v / mv)
        rect(s, x + 124, yy + 4, bw, 14, fill=c)
        text(s, x + 130 + bw, yy, 80, 22, fmt.format(v).replace(".", ",") + unit, size=9, bold=True, color=RED if a == hl else INK)


# ================================================================= 1 COVER
s = prs.slides.add_slide(BL)
pic(s, "s00_cover.jpg", 0, 0, w=W)
rect(s, 0, 0, 760, H, fill="1B1D1E")
text(s, 70, 120, 640, 30, "ÉTUDE DE CAS  ·  CRITIQUE URBAINE", size=11, font=TF, color="E07A3F", spc=400)
text(s, 70, 170, 680, 220, "TANGER,\nVILLE-VITRINE", size=46, font=TF, color=WHITE, spc=300, spacing=1.0)
text(s, 70, 420, 600, 120, "Quand l'urbanisme se met au service du tourisme. Lecture critique de la transformation de la ville, 2007-2030.", size=14, color="D6D2CB", spacing=1.35)
seg(s, 70, 580, 170, 580, "E07A3F", 3)
text(s, 70, 610, 600, 120, "[Nom de l'étudiant·e]\n[Cours / atelier]  ·  [Professeur·e]\n[Date]", size=12, color="D6D2CB", spacing=1.5)
text(s, 70, 900, 620, 80, "Toutes les cartes de ce dossier sont produites à partir de données réelles : OpenStreetMap / Overture Maps, ESA WorldCover, Copernicus DEM, Sentinel-2, HCP.", size=9.5, color="A6A29B", spacing=1.35)
text(s, 1180, 1040, 700, 20, "Image : Sentinel-2 L2A, 19 juillet 2025 (Copernicus / ESA)", size=8, color="D6D2CB", align="r", font=TF, spc=200)

# ================================================================= 2 PROBLEMATIQUE
s = slide(2, "Problématique et méthode", "Une hypothèse, testée par des données spatiales réelles et quatre cas d'étude.",
          kpis=[("6", "jeux de données\nspatiales"), ("4", "cas d'étude"), ("1,28 M", "habitants dans\nla commune (2024)")])
text(s, 40, 180, 1100, 140, [[("« Pour qui construit-on Tanger : pour ceux qui la visitent, ou pour ceux qui y vivent ? »", {"bold": True})]], size=26, font=TF, color=INK, spacing=1.1)
text(s, 40, 320, 1040, 160, [[("Hypothèse. ", {"bold": True, "color": RED}),
     ("Depuis 2007, les grands projets de Tanger (port de plaisance, corniche, médina restaurée, gare TGV, grand stade) obéissent d'abord à une logique d'image et d'attractivité touristique. "
      "La croissance démographique, elle, se fait dans des quartiers périphériques qui ne reçoivent ni la même attention ni les mêmes équipements.", {})]], size=12, spacing=1.45)
ptitle(s, 40, 520, "Démarche")
steps = [("01", "TERRITOIRE", "Le détroit, les pôles, la croissance"), ("02", "VILLE", "Occupation du sol, répartition des usages"), ("03", "CAS", "Port · médina · stade · zones d'emploi"), ("04", "SYNTHÈSE", "Enjeux spatiaux et propositions")]
for i, (n, t, d) in enumerate(steps):
    x = 40 + i * 265
    rect(s, x, 556, 245, 150, fill=WHITE, line=RULE, lw=1)
    text(s, x + 20, 566, 80, 40, n, size=22, font=TF, color=RED, bold=True)
    text(s, x + 20, 622, 210, 24, t, size=12, font=TF, bold=True, spc=300)
    text(s, x + 20, 652, 210, 50, d, size=9.5, color=MUTED, spacing=1.25)
    if i < 3:
        seg(s, x + 247, 631, x + 263, 631, INK, 1.5, head=True)
ptitle(s, 1200, 180, "Données mobilisées")
data = [("OpenStreetMap / Overture Maps", "release 2026-09", "voirie (62 000 tronçons), bâti, 1 674 arrêts de bus, lieux (hôtels, écoles, santé), limites"),
        ("ESA WorldCover 2021", "10 m", "occupation du sol : bâti, arbres, cultures, eau"),
        ("Copernicus DEM GLO-30", "30 m", "relief, ombrage, coupes topographiques"),
        ("Sentinel-2 L2A", "10 m · 2017 / 2025", "images satellite avant / après"),
        ("HCP, RGPH 2014 et 2024", "arrondissements", "population et croissance"),
        ("Presse et rapports publics", "2013-2025", "budgets, Cour des comptes, ONT, FIFA")]
for i, (a, b, c) in enumerate(data):
    y = 216 + i * 82
    seg(s, 1200, y, 1880, y, RULE, 1)
    text(s, 1200, y + 10, 430, 24, a, size=11, bold=True)
    text(s, 1640, y + 12, 240, 24, b, size=9, color=RED, align="r", font=TF, spc=150)
    text(s, 1200, y + 38, 680, 40, c, size=9.5, color=MUTED, spacing=1.2)
text(s, 40, 750, 1080, 200, [[("Précaution. ", {"bold": True}), ("Les lieux (hôtels, écoles, santé) viennent de bases collaboratives et commerciales : elles sont incomplètes, surtout dans les quartiers populaires. "
     "Cette sous-représentation est elle-même un indice de la visibilité inégale des quartiers. Les écarts mesurés sont à lire comme des ordres de grandeur.", {})]], size=10, color=MUTED, spacing=1.4)

# ================================================================= 3 TERRITOIRE (dark)
s = slide(3, "Un territoire tourné vers le détroit", "Les pôles de l'image sur la côte, les pôles productifs le long des autoroutes.",
          kpis=[("1,49 M", "habitants, préfecture\nTanger-Assilah (2024)"), ("+40 %", "en dix ans"), ("167 100", "croisiéristes\nen 2025")], dark=True)
pic(s, "m01_territoire.jpg", 40, 150, w=1250)
x = 1330
ptitle(s, x, 150, "Les grands repères", color="E0E0E0")
poles = [(1, RED, "Port Tanger Ville et Tanja Marina Bay"), (7, RED, "Gare TGV Tanger Ville (Al Boraq, 2018)"), (8, RED, "Grand stade Ibn Batouta (Mondial 2030)"),
         (9, RED, "Cap Spartel, grottes d'Hercule"), (10, RED, "Cap Malabata, front hôtelier"), (2, BLUE, "Port Tanger Med (2007)"), (4, BLUE, "Zone franche de Tanger"),
         (5, BLUE, "Usine Renault Melloussa"), (6, BLUE, "Tanger Automotive City"), (3, "7D8B96", "Aéroport Ibn Battouta")]
for i, (n, c, t) in enumerate(poles):
    badge(s, x + 13, 196 + i * 34, n, c, r=12, size=8.5)
    text(s, x + 36, 186 + i * 34, 520, 24, t, size=10, color="E0E0E0")
y = 550; ptitle(s, x, y, "Familles", color="E0E0E0")
legend(s, x, y + 30, [("dot", RED, "Image, tourisme, événement"), ("dot", BLUE, "Économie productive (port, industrie)"), ("line", "E07A3F", "Autoroutes et voies express"),
                      ("dash", "EDEDED", "Voie ferrée et LGV"), ("line", WHITE, "Commune de Tanger")], dark=True)
text(s, x, 760, 550, 200, [[("Lecture. ", {"bold": True, "color": "E07A3F"}), ("Tous les projets d'image se concentrent sur la façade maritime de la commune. Les pôles d'emploi sont rejetés à 9-32 km, le long des autoroutes, hors de la ville habitée.", {})]],
     size=10.5, color="D6D2CB", spacing=1.4)

# ================================================================= 4 ARRONDISSEMENTS
s = slide(4, "Qui habite Tanger ?", "Population et croissance des quatre arrondissements de la commune, 2014-2024.",
          kpis=[("1 275 428", "habitants, commune\nde Tanger (2024)"), ("44 %", "vivent à\nBeni Makada"), ("+47 %", "croissance de\nBeni Makada")])
pic(s, "m02_arrondissements.jpg", 40, 150, h=860)
x = 1150
ptitle(s, x, 150, "Population 2014 et 2024")
mx = 600000
for i, a in enumerate(["Beni Makada", "Tanger-Médina", "Charf-Mghogha", "Charf-Souani"]):
    y = 190 + i * 98; r = T.loc[a]
    text(s, x, y, 300, 22, a, size=11, bold=True, color=RED if a == "Beni Makada" else INK)
    rect(s, x, y + 28, 520 * r.pop14 / mx, 18, fill=GREY)
    text(s, x + 520 * r.pop14 / mx + 8, y + 26, 160, 22, f"{int(r.pop14):,}".replace(",", " ") + "  (2014)", size=8.5, color=MUTED)
    rect(s, x, y + 50, 520 * r.pop24 / mx, 18, fill=RED if a == "Beni Makada" else "8C8478")
    text(s, x + 520 * r.pop24 / mx + 8, y + 48, 220, 22, f"{int(r.pop24):,}".replace(",", " ") + f"  ({r.growth*100:+.0f} %)", size=9, bold=True)
y = 600
ptitle(s, x, y, "Densité sur le bâti (hab. / km² bâti)")
for i, a in enumerate(["Beni Makada", "Charf-Souani", "Charf-Mghogha", "Tanger-Médina"]):
    yy = y + 30 + i * 30; v = T.loc[a, "dens"]
    text(s, x, yy, 140, 22, a, size=9, bold=(a == "Beni Makada"))
    rect(s, x + 140, yy + 4, 420 * v / 47000, 14, fill=RED if a == "Beni Makada" else GREY)
    text(s, x + 148 + 420 * v / 47000, yy, 100, 22, f"{v:,.0f}".replace(",", " "), size=9, bold=True)
text(s, x, 790, 720, 180, [[("Lecture. ", {"bold": True, "color": RED}), ("Beni Makada a gagné 181 000 habitants en dix ans, plus que toute la croissance des trois autres arrondissements réunis. Avec près de 47 000 habitants par km² bâti, c'est le cœur démographique de Tanger, à 5 km de la mer. Charf-Souani, central, perd des habitants.", {})]], size=10.5, spacing=1.4)

# ================================================================= 5 OCCUPATION
tree_tot = T.tree_km2.sum(); built_tot = T.built_km2.sum()
s = slide(5, "La ville réelle : occupation du sol", "Ce que voit le satellite : bâti, arbres, cultures et relief, à 10 m de résolution.",
          kpis=[(f"{built_tot:.0f} km²", "bâtis dans la\ncommune"), (f"{tree_tot:.0f} km²", "de couvert\narboré"), (f"{T.loc['Tanger-Médina','tree_km2']/tree_tot*100:.0f} %", "des arbres dans\nTanger-Médina"), (f"{T.loc['Beni Makada','tree_km2']:.1f} km²", "d'arbres à\nBeni Makada")])
pic(s, "m03_occupation.jpg", 40, 150, w=1250)
x = 1330
ptitle(s, x, 150, "Légende")
y = legend(s, x, 182, [("box", SALMON, "Surfaces bâties"), ("box", "A3B78F", "Arbres, forêts (Rmilat, Perdicaris, Malabata)"), ("box", "C8CFAE", "Arbustes"), ("box", "DAD9C2", "Prairies"),
                       ("box", "E5DEC8", "Cultures"), ("box", "A9C3CC", "Mer et plans d'eau"), ("line", RED, "Autoroute, voie express"), ("line", "D9824F", "Route principale"),
                       ("dash", INK, "Voie ferrée"), ("line", INK, "Commune et arrondissements")])
text(s, x, y + 30, 550, 300, [[("Lecture. ", {"bold": True, "color": RED}), ("La ville forme une tache bâtie continue qui s'étire vers le sud et le sud-ouest, le long des voies express. Au nord, le massif forestier de Rmilat et du cap Spartel reste le grand poumon vert, mais il est presque entièrement dans l'arrondissement côtier. "
     "Beni Makada, le plus peuplé, n'a presque plus d'arbres.", {})]], size=10.5, spacing=1.4)

# ================================================================= 6 SATELLITE 2017 / 2025
s = slide(6, "2017 → 2025 : ce que montre le satellite", "Deux images Sentinel-2 prises en juillet, à huit ans d'écart.",
          kpis=[("8 ans", "d'écart"), ("10 m", "résolution\nSentinel-2")])
iw = 910
w1, h1 = pic(s, "s01_ville_2017.jpg", 40, 170, w=iw)
pic(s, "s01_ville_2025.jpg", 40 + iw + 20, 170, w=iw)
text(s, 50, 180, 200, 30, "16 / 07 / 2017", size=14, font=TF, bold=True, color=WHITE, spc=200)
text(s, 60 + iw + 20 - 10, 180, 200, 30, "19 / 07 / 2025", size=14, font=TF, bold=True, color=WHITE, spc=200)
def sat_xy(lon, lat, x0, y0, w, h, box=(-5.955, 35.705, -5.735, 35.815)):
    return x0 + (lon - box[0]) / (box[2] - box[0]) * w, y0 + (box[3] - lat) / (box[3] - box[1]) * h
ann = [(1, -5.8045, 35.7895, "Port reconverti et marina"), (2, -5.8581, 35.7412, "Grand stade : nouvelle toiture"), (3, -5.9053, 35.7155, "Zone franche étendue"), (4, -5.7700, 35.7330, "Lotissements de Mghogha")]
for n, lon, lat, t in ann:
    cx, cy = sat_xy(lon, lat, 40 + iw + 20, 170, iw, h1)
    oval(s, cx, cy, 22, fill=None, line=WHITE, lw=2)
    badge(s, cx + 22, cy - 22, n, RED, r=12, size=8.5)
y = 170 + h1 + 24
for i, (n, _, _, t) in enumerate(ann):
    badge(s, 40 + iw + 20 + 14 + (i % 2) * 450, y + 14 + (i // 2) * 34, n, RED, r=12, size=8.5)
    text(s, 40 + iw + 20 + 36 + (i % 2) * 450, y + 3 + (i // 2) * 34, 400, 24, t, size=10.5)
text(s, 40, y, iw, 120, [[("Lecture. ", {"bold": True, "color": RED}), ("Les transformations les plus visibles depuis l'espace sont les projets d'image : le front portuaire et le grand stade. Pendant ce temps, la ville s'étale au sud par lotissements successifs, sans transport en site propre pour les relier.", {})]], size=10.5, spacing=1.4)

# ================================================================= 7 DEUX VILLES
s = slide(7, "Deux villes dans la ville", "Densité des lieux recensés : à gauche, l'offre pour le visiteur ; à droite, les équipements du quotidien.",
          kpis=[("77 %", "des hébergements à\nmoins de 1 km de la mer"), ("14 %", "du bâti à moins\nde 1 km de la mer"), ("27 %", "des écoles à moins\nde 1 km de la mer")])
mw = 900
wa, ha = pic(s, "m04a_vitrine.jpg", 40, 190, w=mw)
pic(s, "m04b_vecue.jpg", 40 + mw + 40, 190, w=mw)
ptitle(s, 40, 160, "La ville du visiteur : hébergements, sites patrimoniaux, plages", RED, w=900)
ptitle(s, 80 + mw, 160, "La ville de l'habitant : écoles et équipements de santé", BLUE, w=900)
y = 190 + ha + 20
text(s, 40, y, mw, 120, [[("Hébergements : ", {"bold": True}), ("207 dans la commune, 62 % à moins de 500 m de la mer, concentrés entre médina, port et corniche.", {})]], size=10.5, spacing=1.35)
text(s, 80 + mw, y, mw, 120, [[("Écoles et santé : ", {"bold": True}), ("573 lieux, plus diffus, mais leur densité chute fortement dans Beni Makada, pourtant le plus peuplé.", {})]], size=10.5, spacing=1.35)
note_y = y + 80
text(s, 40, note_y, 1840, 60, "Estimation par noyau (KDE) des lieux Overture Places 2026. Distance à la mer calculée sur le trait de côte ESA WorldCover. Calculs de l'auteur.", size=8.5, italic=True, color=MUTED)

# ================================================================= 8 INDICATEURS
s = slide(8, "Indicateurs par arrondissement", "Équipements et services rapportés à la population 2024. Beni Makada en rouge.",
          kpis=[("× 29", "hébergements / hab.\nMédina vs Beni Makada"), ("× 6", "équipements de santé\n/ hab."), ("× 4", "écoles / hab.")])
cols = [(40, 230), (660, 230), (1280, 230), (40, 560), (660, 560), (1280, 560)]
specs = [("Hébergements touristiques / 10 000 hab.", "lodging_10k", "{:.1f}", ""), ("Écoles / 10 000 hab.", "school_10k", "{:.1f}", ""),
         ("Équipements de santé / 10 000 hab.", "health_10k", "{:.1f}", ""), ("Arrêts de bus / 10 000 hab.", "stops_10k", "{:.1f}", ""),
         ("Parcs urbains (m² / habitant)", "upark_m2_hab", "{:.2f}", " m²"), ("Couvert arboré (km²)", "tree_km2", "{:.1f}", " km²")]
for (x, y), (t, k, f, u) in zip(cols, specs):
    rect(s, x, y - 20, 590, 290, fill=WHITE, line=RULE, lw=1)
    hbars(s, x + 24, y, 560, t, k, f, u)
text(s, 40, 880, 1840, 120, [[("Lecture. ", {"bold": True, "color": RED}), ("À chaque indicateur, l'arrondissement le plus peuplé arrive dernier ou presque. Un habitant de Beni Makada dispose de 0,17 m² de parc, de six fois moins d'équipements de santé recensés et de trois fois moins d'arrêts de bus qu'un habitant de Tanger-Médina. "
     "Les bases de lieux étant moins complètes dans les quartiers populaires, ces écarts sont à lire comme des ordres de grandeur.", {})]], size=10.5, spacing=1.4)

# ================================================================= 9 LITTORAL
s = slide(9, "Le littoral capté", "Bande de 1 km depuis le trait de côte : le bâti qui s'y trouve et les hébergements recensés.",
          kpis=[("62 %", "des hébergements\nà moins de 500 m"), ("14 %", "du bâti dans\nla bande de 1 km"), ("9 %", "des écoles\nà moins de 500 m")])
pic(s, "m05_littoral.jpg", 40, 160, w=1250)
x = 1330
ptitle(s, x, 160, "Légende")
y = legend(s, x, 192, [("box", "B7553A", "Bâti dans la bande littorale de 1 km"), ("box", "9A9184", "Bâti au-delà de 1 km"), ("dash", RED, "Limite de la bande de 1 km"),
                       ("dot", RED, "Hébergement touristique"), ("sq", BLUE, "École ou équipement de santé"), ("line", RED, "Voie express"), ("line", INK, "Commune et arrondissements")])
text(s, x, y + 30, 550, 400, [[("Lecture. ", {"bold": True, "color": RED}), ("La bande littorale ne porte qu'un septième du bâti de la commune, mais plus des trois quarts des hébergements. La côte est traitée comme un produit : marina, hôtels de Malabata, résidences du cap Spartel. "
     "La majorité des habitants vit à plus de 3 km de la mer et n'y accède que par la route.", {})], "",
     [("Médiane de la distance à la mer : ", {"bold": True}), ("390 m pour un hébergement, 1,8 km pour une école, 4,9 km pour le bâti de Beni Makada.", {})]], size=10.5, spacing=1.4)

ptitle(s, 40, 860, "Distance médiane à la mer")
for i, (lab, v, c) in enumerate([("Hébergements", 390, RED), ("Équipements de santé", 1436, BLUE), ("Écoles", 1815, BLUE), ("Bâti de Beni Makada", 4927, GREY)]):
    xx = 40 + (i % 2) * 640; yy = 892 + (i // 2) * 44
    text(s, xx, yy, 200, 22, lab, size=9.5)
    rect(s, xx + 200, yy + 4, 330 * v / 4927, 16, fill=c)
    text(s, xx + 208 + 330 * v / 4927, yy, 100, 22, (f"{v} m" if v < 1000 else f"{v/1000:.1f} km".replace(".", ",")), size=9.5, bold=True)

# ================================================================= 10 PORT
s = slide(10, "Cas 1 · Le port : du travail au loisir", "Le front portuaire vu par Sentinel-2, avant et après la reconversion.",
          kpis=[("6,2 Md", "MAD, programme\nde reconversion"), ("1 400", "anneaux de\nplaisance prévus"), ("7 / 6 / 2018", "inauguration")])
pw = 610
w1, h1 = pic(s, "s02_port_2017.jpg", 40, 170, w=pw)
pic(s, "s02_port_2025.jpg", 40 + pw + 20, 170, w=pw)
text(s, 52, 180, 200, 30, "JUILLET 2017", size=13, font=TF, bold=True, color=WHITE, spc=200)
text(s, 52 + pw + 20, 180, 200, 30, "JUILLET 2025", size=13, font=TF, bold=True, color=WHITE, spc=200)
y = 170 + h1 + 20
text(s, 40, y, pw, 120, "Chantier en cours : la ville tourne encore le dos à un port de travail, de ferries et de pêche.", size=10.5, spacing=1.35, color=MUTED)
text(s, 60 + pw, y, pw, 120, "Bassins de plaisance, esplanades et parkings : le front d'eau est devenu un espace de loisir et de consommation.", size=10.5, spacing=1.35, color=MUTED)
x = 1330
ptitle(s, x, 170, "Chiffres clés")
for i, (v, d) in enumerate([("6,2 Md MAD", "programme de reconversion du port"), ("635 M MAD", "coût de la marina Tanja Marina Bay"), ("1,19 Md MAD", "coût du nouveau port de pêche"), ("1 400", "anneaux prévus : plus grande marina du Maroc")]):
    text(s, x, 200 + i * 76, 550, 36, v, size=19, font=TF, bold=True, color=RED if i < 2 else INK)
    text(s, x, 244 + i * 76, 550, 24, d, size=9.5, color=MUTED)
seg(s, x, 510, 1880, 510, RULE, 1)
ptitle(s, x, 526, "Bilan")
for i, (sg, t, c) in enumerate([("+", "Promenade publique continue sur la baie", GREEN), ("−", "Le travail portuaire sort du centre", RED), ("−", "Usages payants : plaisanciers, croisiéristes", RED),
                                 ("−", "La route littorale sépare toujours ville et plage", RED)]):
    badge(s, x + 12, 572 + i * 40, sg, c, r=12, size=11)
    text(s, x + 36, 561 + i * 40, 520, 24, t, size=10.5)
text(s, x, 750, 550, 220, [[("Lecture. ", {"bold": True, "color": RED}), ("Le port n'a pas été rendu aux Tangérois : il a été converti en vitrine. Le programme visait à « réconcilier la ville et la mer » ; le résultat est un espace calibré pour le visiteur et le plaisancier.", {})]], size=10.5, spacing=1.4)

# ================================================================= 11 COUPES
s = slide(11, "Coupes : la pente et la distance", "Profils topographiques réels (Copernicus DEM), bâti relevé par satellite le long des tracés.",
          kpis=[("90 m", "de dénivelé entre\nMarshan et le port"), ("3 km", "coupe A-A'"), ("4,7 km", "coupe B-B'")])
ptitle(s, 40, 160, "Coupe A-A' : Marshan, Kasbah, médina, port", RED)
pic(s, "m09_coupeA.jpg", 40, 190, w=1330)
ptitle(s, 40, 590, "Coupe B-B' : corniche, ville nouvelle, Beni Makada", RED)
wB, hB = pic(s, "m09_coupeB.jpg", 40, 616, w=1330)
pic(s, "m09_coupe_loc.jpg", 1450, 186, w=430)
text(s, 1450, 680, 430, 300, [[("Lecture. ", {"bold": True, "color": RED}), ("La vieille ville domine la mer, mais n'y descend que par des escaliers, des portes et une route à plusieurs voies. Depuis Beni Makada, la mer est à près de 5 km, derrière les collines : la corniche n'est pas un espace de proximité pour la majorité des habitants.", {})]], size=10.5, spacing=1.4)

# ================================================================= 12 MEDINA
s = slide(12, "Cas 2 · La médina, vitrine patrimoniale", "Bâti, voirie, remparts et lieux recensés dans la médina et la Kasbah.",
          kpis=[("65", "hébergements recensés\ndans 0,9 km²"), ("850 M", "MAD, programme\n2020-2024"), ("328", "bâtiments menaçant\nruine")])
pic(s, "m07_medina.jpg", 40, 150, h=860)
x = 1230
ptitle(s, x, 150, "Légende")
y = legend(s, x, 182, [("dot", RED, "Hébergement touristique (hôtel, riad, maison d'hôtes)"), ("diamond", GOLD, "Site patrimonial, musée, monument"), ("box", "B9B1A5", "Bâti relevé (OSM)"),
                       ("line", INK, "Remparts"), ("box", "C9D3BB", "Jardins et espaces verts"), ("box", "A9C3CC", "Mer et bassins portuaires")], w=640)
ptitle(s, x, y + 30, "Où va l'argent ?")
for i, (lab, v, c) in enumerate([("Programme médina 2020-2024", 850, RED), ("Bâtiments menaçant ruine 2024-2027", 98, BLUE)]):
    yy = y + 64 + i * 64
    text(s, x, yy, 500, 22, lab, size=10, bold=True)
    rect(s, x, yy + 26, 480 * v / 850, 22, fill=c)
    text(s, x + 480 * v / 850 + 10, yy + 24, 160, 24, f"{v} M MAD", size=11, bold=True, font=TF, color=c)
text(s, x, y + 210, 640, 300, [[("Lecture. ", {"bold": True, "color": RED}), ("Dans moins d'un kilomètre carré, on recense 65 hébergements touristiques pour 5 équipements de santé. Les travaux suivent les portes, places et parcours des visiteurs, alors que 328 bâtiments habités menacent ruine et relèvent d'un budget près de neuf fois plus petit. "
     "La médina se transforme en décor : logements convertis en maisons d'hôtes, commerces tournés vers le visiteur.", {})]], size=10.5, spacing=1.4)

# ================================================================= 13 MONDIAL
s = slide(13, "Cas 3 · Mondial 2030 : l'urbanisme d'exception", "Le grand stade Ibn Batouta vu par Sentinel-2, et l'évaluation de la ville par la FIFA.",
          kpis=[("4,0 / 5", "stade"), ("2,6 / 5", "transports"), ("2,2 / 5", "hôtels")])
sw = 600
w1, h1 = pic(s, "s03_stade_2017.jpg", 40, 170, w=sw)
pic(s, "s03_stade_2025.jpg", 60 + sw, 170, w=sw)
text(s, 52, 180, 200, 30, "JUILLET 2017", size=13, font=TF, bold=True, color=WHITE, spc=200)
text(s, 72 + sw, 180, 200, 30, "JUILLET 2025", size=13, font=TF, bold=True, color=WHITE, spc=200)
y = 170 + h1 + 30
ptitle(s, 40, y, "Deux chemins pour construire")
for row, (lab, c, items, fill) in enumerate([("LA RÈGLE", BLUE, ["Plan d'aménagement", "Enquête publique", "Avis des habitants", "Permis"], None),
                                             ("L'EXCEPTION", RED, ["Événement (2030)", "Urgence « date FIFA »", "Autorisation exceptionnelle", "Hôtels, résidences"], RED)]):
    yy = y + 34 + row * 90
    text(s, 40, yy + 18, 118, 30, lab, size=10, font=TF, bold=True, color=c, spc=100)
    for k, t in enumerate(items):
        bx = 160 + k * 270
        rect(s, bx, yy, 240, 60, fill=fill or WHITE, line=c, lw=1.5)
        text(s, bx + 8, yy, 224, 60, t, size=10.5, bold=True, color=WHITE if fill else c, align="c", anchor="m")
        if k < 3: seg(s, bx + 242, yy + 30, bx + 268, yy + 30, c, 1.5, head=True)
x = 1330
ptitle(s, x, 170, "Capacité du stade (places)")
for i, (lab, v, c) in enumerate([("Origine", 45000, GREY), ("Après extension", 65000, RED), ("Objectif Mondial", 85000, "E4A08A")]):
    yy = 206 + i * 50
    text(s, x, yy, 160, 22, lab, size=10)
    rect(s, x + 160, yy + 2, 300 * v / 85000, 20, fill=c)
    text(s, x + 170 + 300 * v / 85000, yy, 100, 22, (">40 000" if i == 0 else f"{v:,}".replace(",", " ")), size=10, bold=True)
text(s, x, 380, 550, 600, [[("Lecture. ", {"bold": True, "color": RED}), ("Le stade a été agrandi deux fois et sa toiture est visible depuis l'espace. La FIFA juge pourtant les transports (2,6/5) et les hôtels (2,2/5) insuffisants. "
     "La réponse annoncée : des permis hôteliers exceptionnels, pour tenir une date fixée de l'extérieur. Ce qui peut être livré à temps (stade, hôtels, façades) passe devant ce qui demande du temps : logement, transport, équipements de quartier.", {})], "",
     [("L'héritage du Mondial devra se mesurer dans les quartiers, pas dans les tribunes.", {"bold": True})]], size=10.5, spacing=1.4)

# ================================================================= 14 INDUSTRIE / CHRAFATE
s = slide(14, "Cas 4 · Les ouvriers loin de leurs emplois", "Distances entre le quartier le plus peuplé et les zones d'emploi industrielles.",
          kpis=[("9-32 km", "entre Beni Makada et\nles zones d'emploi"), ("150 000", "habitants prévus\nà Chrafate (2009)"), ("36", "habitants comptés\n(Cour des comptes, 2024)")])
pic(s, "m12_industrie.jpg", 40, 150, w=1250)
x = 1330
ptitle(s, x, 150, "Légende")
y = legend(s, x, 182, [("box", BLUE, "Zones industrielles (OSM)"), ("box", "BFB3A3", "Bâti"), ("dash", RED, "Distance à vol d'oiseau depuis Beni Makada"), ("line", RED, "Autoroute, voie express"), ("dash", INK, "Voie ferrée")])
ptitle(s, x, y + 26, "Chrafate : prévu vs compté")
for i in range(100):
    cx = x + 12 + (i % 20) * 27; cy = y + 70 + (i // 20) * 27
    oval(s, cx, cy, 9, fill="C9D5E3", line=BLUE, lw=0.8)
text(s, x, y + 210, 550, 30, "chaque cercle = 1 500 habitants prévus ; les 36 habitants comptés tiennent dans moins d'un cercle.", size=9, italic=True, color=MUTED)
text(s, x, y + 260, 550, 300, [[("Lecture. ", {"bold": True, "color": RED}), ("La ville nouvelle de Chrafate, lancée en 2009 près des usines pour loger 150 000 personnes, ne comptait que 36 habitants selon la Cour des comptes. Les ouvriers vivent donc dans les quartiers denses de Tanger et font chaque jour 9 à 32 km. "
     "Les projets pour le visiteur ont été livrés en moins de dix ans ; celui pour les travailleurs attend depuis 2009.", {})]], size=10.5, spacing=1.4)

# ================================================================= 15 MOBILITE
s = slide(15, "Mobilité : le bâti hors de portée du bus", "Surfaces bâties à moins ou à plus de 300 m d'un arrêt de bus recensé.",
          kpis=[("1 674", "arrêts de bus\nrecensés"), ("31 %", "du bâti à plus de\n300 m d'un arrêt"), ("0", "ligne de tram\nou de BHNS")])
pic(s, "m06_mobilite.jpg", 40, 150, w=1250)
x = 1330
ptitle(s, x, 150, "Légende")
y = legend(s, x, 182, [("box", "A9B6C6", "Bâti desservi (moins de 300 m d'un arrêt)"), ("box", RED, "Bâti non desservi (plus de 300 m)"), ("dot", BLUE, "Arrêt de bus"),
                       ("line", RED, "Autoroute, voie express"), ("dash", INK, "Voie ferrée et LGV"), ("dot", RED, "Gare TGV Tanger Ville")])
hbars(s, x, y + 30, 550, "Part du bâti desservi", "served", "{:.0%}", maxv=1)
text(s, x, y + 200, 550, 400, [[("Lecture. ", {"bold": True, "color": RED}), ("Près d'un tiers des surfaces bâties de la commune est à plus de 300 m d'un arrêt, surtout dans les extensions récentes du sud et de Gzenaya. Tanger est reliée à Casablanca en 2 h 10 par le TGV, mais ses quartiers ne sont reliés entre eux que par le bus, le grand taxi et la voiture.", {})]], size=10.5, spacing=1.4)

# ================================================================= 16 SYNTHESE
s = slide(16, "Synthèse spatiale et enjeux", "Trois figures superposées : la vitrine littorale, la périphérie dense, la lisière forestière sous pression.",
          kpis=[("1 km", "de bande littorale\nconcentrant l'image"), ("567 000", "habitants dans la\npériphérie dense"), ("5", "enjeux")])
pic(s, "m10_synthese.jpg", 40, 150, w=1250)
x = 1330
ptitle(s, x, 150, "Légende")
y = legend(s, x, 182, [("box", "E9B6A8", "Vitrine littorale : bande de 1 km"), ("hatch", BLUE, "Périphérie dense sous-équipée (Beni Makada)"), ("box", "2F5E2A", "Lisière forestière au contact du bâti"),
                       ("box", "C6B9A8", "Bâti"), ("box", "B4C6A2", "Couvert arboré")])
ptitle(s, x, y + 26, "Enjeux")
enj = [("Front de mer et port", "la vitrine concentre l'investissement et l'hébergement"), ("Beni Makada", "44 % des habitants, le moins d'équipements et d'arbres"),
       ("Rmilat, Perdicaris, cap Spartel", "la forêt devient la prochaine réserve foncière"), ("Grand stade", "un investissement d'événement, loin des transports"),
       ("Gzenaya, zone franche", "l'emploi sans la ville")]
for i, (t, d) in enumerate(enj):
    yy = y + 60 + i * 58
    badge(s, x + 13, yy + 12, i + 1, INK, r=13, size=9)
    text(s, x + 38, yy, 520, 24, t, size=11, bold=True)
    text(s, x + 38, yy + 24, 520, 24, d, size=9.5, color=MUTED)

# ================================================================= 17 PROPOSITIONS
s = slide(17, "Propositions : retourner la vitrine", "Cinq interventions localisées, et une règle : juger chaque projet à ce qu'il apporte au quartier le plus éloigné.",
          kpis=[("17 km", "de ligne structurante\nGzenaya-gare"), ("6", "équipements de\nquartier prioritaires"), ("328", "logements à sauver\nen médina")])
pic(s, "m11_propositions.jpg", 40, 150, w=1250)
x = 1330
props = [("Reconnecter", "Liaison piétonne continue médina, port, corniche ; apaisement de la route littorale.", RED),
         ("Habiter la médina", "Logement d'abord : les 328 bâtiments en péril avant les façades ; encadrer la location de courte durée.", RED),
         ("Relier les quartiers", "Une ligne de tram ou de BHNS de 17 km, tracée sur la voirie existante : Gzenaya, Mesnana, Beni Makada, centre, gare.", INK),
         ("Protéger la lisière", "Limite verte ferme autour de Rmilat, Perdicaris et cap Spartel ; rivage public garanti.", GREEN),
         ("Équiper la périphérie", "Héritage Mondial : chaque dirham de l'événement couplé à une école, un centre de santé ou un parc à Beni Makada et Mghogha.", GOLD)]
for i, (t, d, c) in enumerate(props):
    yy = 160 + i * 150
    badge(s, x + 14, yy + 14, i + 1, INK, r=14, size=10)
    rect(s, x + 40, yy + 4, 6, 22, fill=c)
    text(s, x + 56, yy + 2, 500, 28, t.upper(), size=13, font=TF, bold=True, spc=200)
    text(s, x + 40, yy + 38, 510, 100, d, size=10.5, spacing=1.35)
y = 920
legend(s, x, y, [("line", INK, "Ligne structurante proposée, tracée sur la voirie réelle"), ("sq", GOLD, "Équipement de quartier prioritaire")])

# ================================================================= 18 CONCLUSION
s = slide(18, "Conclusion", None)
text(s, 40, 130, 1840, 140, [[("Tanger a été refaite pour être vue. ", {"bold": True}), ("Les données le montrent : la vitrine se mesure en mètres depuis la mer, la ville vécue en kilomètres.", {"color": RED, "bold": True})]], size=26, font=TF, spacing=1.15)
for k, (v, d, c) in enumerate([("77 % / 14 %", "des hébergements contre part du bâti à moins de 1 km de la mer", RED), ("44 %", "des habitants à Beni Makada, dernier pour presque tous les équipements", BLUE),
                                ("0,17 m²", "de parc par habitant à Beni Makada", BLUE), ("36 / 150 000", "habitants comptés vs prévus à Chrafate", BLUE), ("850 vs 98 M", "MAD : décor de la médina vs bâtiments en ruine", RED)]):
    x = 40 + k * 372
    rect(s, x, 300, 352, 190, fill=WHITE, line=RULE, lw=1)
    rect(s, x, 300, 352, 6, fill=c)
    text(s, x + 22, 326, 320, 50, v, size=26, font=TF, bold=True, color=c)
    text(s, x + 22, 384, 310, 100, d, size=10, spacing=1.3)
text(s, 40, 520, 1840, 50, "« Une ville se juge moins à son front de mer qu'à la vie quotidienne de son quartier le plus éloigné. »", size=17, italic=True, align="c", font=TF)
seg(s, 40, 590, 1880, 590, RULE, 1)
ptitle(s, 40, 606, "Sources")
src = ["Données spatiales : OpenStreetMap contributors via Overture Maps Foundation, release 2026-09-23 (voirie, bâti, lieux, infrastructures, limites administratives) · ESA WorldCover 2021 v200 (10 m) · Copernicus DEM GLO-30 · Sentinel-2 L2A, scènes du 16/07/2017 et du 19/07/2025 (Copernicus, via AWS Open Data).",
       "Population : HCP, Recensement général de la population et de l'habitat 2014 et 2024, données par arrondissement (via citypopulation.de).",
       "Tourisme : Observatoire national du tourisme via Hespress FR (nov. 2025) ; port de Tanger Ville via Finances News Hebdo (2025) et Le360.",
       "Projets : Ministère de l'Économie et des Finances (2018), Médias24 (2013-2014, 2018, 2020) ; Le360, Finances News (juin 2018) ; TelQuel (2024) ; Le Desk (2025) citant la Cour des comptes (2024) ; Medias24 (2024) et Bladi (2024) pour le stade et la FIFA.",
       "Méthode : reprojection UTM 30N ; densités par estimation par noyau ; distances euclidiennes au trait de côte ; desserte = 300 m autour des arrêts ; profils extraits du MNT. Tous les calculs et cartes sont de l'auteur."]
text(s, 40, 636, 1840, 380, src, size=9.5, color=MUTED, spacing=1.45)

for sl in prs.slides:
    for st in list(sl.shapes._spTree.iter(qn("p:style"))):
        st.getparent().remove(st)
prs.save("Tanger_ville_vitrine_planches.pptx")
print("ok", len(prs.slides))
